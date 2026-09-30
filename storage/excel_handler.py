"""
Storage and export handler for HomeDesk Facility Management Portal.
PostgreSQL is the primary production data store.
Excel is used strictly for on-demand reporting and data export.
"""
from datetime import datetime
import io
import logging
from pathlib import Path
from typing import Optional, Dict, Any
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter

from storage.config import EXCEL_FILE
from database.repository import (
    init_database,
    save_lead_to_db,
    save_requirement_to_db,
    get_latest_lead_by_mobile,
    get_all_requirements_for_export,
    get_all_leads_for_export,
)

logger = logging.getLogger(__name__)

Sheets = {
    "All Leads": [
        "Lead ID",
        "Date Time",
        "Name",
        "Mobile",
        "Service Type",
        "City",
        "Budget",
        "Status",
    ],
    "Cook": [
        "Lead ID",
        "Date Time",
        "Name",
        "Mobile",
        "Email",
        "Address",
        "City",
        "State",
        "Pincode",
        "Start Date",
        "Preferred Timing",
        "Budget",
        "Cuisine Type",
        "Meals Per Day",
        "Additional Notes",
        "Status",
    ],
    "Driver": [
        "Lead ID",
        "Date Time",
        "Name",
        "Mobile",
        "Email",
        "Address",
        "City",
        "State",
        "Pincode",
        "Start Date",
        "Preferred Timing",
        "Budget",
        "Vehicle Type",
        "License Required",
        "Additional Notes",
        "Status",
    ],
    "Security Guard": [
        "Lead ID",
        "Date Time",
        "Name",
        "Mobile",
        "Email",
        "Address",
        "City",
        "State",
        "Pincode",
        "Start Date",
        "Preferred Timing",
        "Budget",
        "Day/Night Shift",
        "Residential/Commercial",
        "Additional Notes",
        "Status",
    ],
}


def auto_adjust_columns(ws):
    """Adjusts column widths dynamically based on content."""
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = max(max_length + 3, 12)


def create_excel_if_not_exists():
    """Initializes the database schema if needed."""
    try:
        init_database()
    except Exception as exc:
        logger.warning("Database init check in excel_handler warning: %s", exc)


def save_lead(
    service_name: str,
    lead_data: Dict[str, Any],
    organization_id: Optional[int] = None,
) -> Dict[str, Any]:
    """
    Saves a requirement to PostgreSQL as the primary production data store.
    No longer performs synchronous file writes to local Excel spreadsheets.
    """
    lead_data.setdefault("Date Time", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    lead_data.setdefault("Status", "New")

    try:
        saved = save_requirement_to_db(service_name, lead_data, organization_id=organization_id)
        lead_data.update(saved)
        return lead_data
    except Exception as exc:
        logger.error("PostgreSQL save failed: %s", exc)
        raise


def find_latest_request_by_mobile(
    mobile: str,
    organization_id: Optional[int] = None,
) -> Optional[Dict[str, Any]]:
    """
    Retrieves the latest request for a mobile number from PostgreSQL.
    """
    mobile_clean = str(mobile or "").strip()
    if not mobile_clean:
        return None

    try:
        db_lead = get_latest_lead_by_mobile(mobile_clean, organization_id=organization_id)
        if db_lead:
            return db_lead
    except Exception as exc:
        logger.error("Database lookup error for mobile '%s': %s", mobile_clean, exc)

    return None


def get_excel_export_bytes(
    organization_id: Optional[int] = None,
    category_id: Optional[int] = None,
) -> bytes:
    """
    Generates an Excel workbook dynamically in-memory from PostgreSQL data
    and returns raw bytes for browser download.
    If scoped to an organization/category, includes only the sheets matching that category.
    """
    all_data = get_all_requirements_for_export(
        organization_id=organization_id,
        category_id=category_id,
    )
    wb = Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    allowed_sheet_names = None
    if organization_id or category_id:
        target_cat_id = category_id
        if target_cat_id is None and organization_id:
            try:
                from database.connection import get_db
                from database.models import Organization
                with get_db() as db:
                    org = db.query(Organization).filter_by(id=organization_id).first()
                    if org:
                        target_cat_id = org.category_id
            except Exception:
                pass

        if target_cat_id:
            try:
                from database.connection import get_db
                from database.models import Category
                from database.repository import CATEGORY_NAME_TO_SERVICE
                with get_db() as db:
                    cat = db.query(Category).filter_by(id=target_cat_id).first()
                    if cat:
                        display = cat.display_name or CATEGORY_NAME_TO_SERVICE.get(cat.name, cat.name)
                        allowed_sheet_names = [display]
            except Exception:
                pass

    sheets_to_export = [
        (s, headers) for s, headers in Sheets.items()
        if allowed_sheet_names is None or s in allowed_sheet_names
    ]
    if not sheets_to_export:
        sheets_to_export = list(Sheets.items())

    for sheet_name, headers in sheets_to_export:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(headers)

        rows = all_data.get(sheet_name, [])
        for item in rows:
            row = [item.get(h, "") for h in headers]
            ws.append(row)

        auto_adjust_columns(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    buffer.seek(0)
    return buffer.getvalue()


def export_leads_to_excel(
    organization_id: Optional[int] = None,
    category_id: Optional[int] = None,
    target_path: Optional[Path] = None,
) -> Path:
    """
    Reads requirements from PostgreSQL and generates an Excel workbook on demand.
    """
    export_file = Path(target_path or EXCEL_FILE)
    export_file.parent.mkdir(parents=True, exist_ok=True)

    excel_bytes = get_excel_export_bytes(
        organization_id=organization_id,
        category_id=category_id,
    )
    with open(export_file, "wb") as f:
        f.write(excel_bytes)

    return export_file