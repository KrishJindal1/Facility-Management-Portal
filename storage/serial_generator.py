import logging
from openpyxl import load_workbook
from storage.config import EXCEL_FILE
from database.repository import generate_next_lead_id

logger = logging.getLogger(__name__)


def generate_serial(service_name, organization_id=None):
    """
    Generates a sequential Lead ID (e.g. Cook-001, Driver-001, Security-001).
    Primary source is PostgreSQL scoped to the tenant.
    """
    try:
        return generate_next_lead_id(service_name, organization_id)
    except Exception as exc:
        logger.warning("Database serial generation failed, falling back to Excel: %s", exc)


    if not EXCEL_FILE.exists():
        prefix_map = {
            "Cook": "Cook",
            "Driver": "Driver",
            "Security Guard": "Security",
        }
        return f"{prefix_map.get(service_name, service_name)}-001"

    wb = load_workbook(EXCEL_FILE)
    ws = wb[service_name]

    prefix_map = {
        "Cook": "Cook",
        "Driver": "Driver",
        "Security Guard": "Security",
    }
    prefix = prefix_map[service_name]
    used_numbers = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        lead_id = row[0]
        if not lead_id:
            continue
        try:
            number = int(str(lead_id).split("-")[-1])
            used_numbers.add(number)
        except (ValueError, IndexError):
            continue

    wb.close()

    serial_number = 1
    while serial_number in used_numbers:
        serial_number += 1

    return f"{prefix}-{serial_number:03d}"