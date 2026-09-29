from openpyxl import load_workbook
from config import EXCEL_FILE

wb = load_workbook(EXCEL_FILE)


for sheet in wb.sheetnames:
    print(f"\n--- {sheet} ---")

    ws = wb[sheet]

    for row in ws.iter_rows(values_only=True):
        print(row)