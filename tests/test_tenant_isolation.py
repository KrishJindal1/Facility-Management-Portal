"""
Automated Category-Based Tenant Isolation Test Suite.
Verifies Phase 2 business model and isolation rules:
1. Normal user submits requirements belonging to Category and User.
2. Organization registers with one category and sees ONLY requirements matching its category.
3. Cook Organization (Org 1) sees Cook requirements, cannot see Driver or Security requirements.
4. Driver Organization (Org 2) sees Driver requirements, cannot see Cook or Security requirements.
5. Cross-organization lookups by mobile or Lead ID are strictly prevented across categories.
6. Data access layer enforces category validation and bounds.
7. Excel export generated for an organization only contains its category's requirements.
"""
import io
import unittest
from openpyxl import load_workbook
from database.connection import get_db
from database.models import Organization, User, Requirement, Category
from database.repository import (
    save_requirement_to_db,
    save_lead_to_db,
    get_latest_lead_by_mobile,
    get_lead_by_id,
    generate_next_lead_id,
    get_all_requirements_for_export,
    _resolve_tenant_id,
    init_database,
)
from storage.excel_handler import get_excel_export_bytes


class TestTenantIsolation(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        init_database()

    def setUp(self):
        with get_db() as db:
            org1 = db.query(Organization).filter_by(slug="homedesk").first()
            org2 = db.query(Organization).filter_by(slug="acme").first()
            assert org1 is not None, "Tenant 1 (homedesk) must exist"
            assert org2 is not None, "Tenant 2 (acme) must exist"
            self.org1_id = int(org1.id)
            self.org2_id = int(org2.id)
            self.org1_cat_id = int(org1.category_id)
            self.org2_cat_id = int(org2.category_id)

    def test_01_multiple_users_per_organization(self):
        """Requirement: An organization can have multiple staff/admin users."""
        with get_db() as db:
            for em in ("alice@homedesk.com", "bob@homedesk.com", "charlie@acme.com"):
                old = db.query(User).filter_by(email=em).first()
                if old:
                    db.delete(old)
            db.flush()

            u1 = User(
                organization_id=self.org1_id,
                name="Alice Manager",
                phone="9800000001",
                email="alice@homedesk.com",
                role="organization",
            )
            u2 = User(
                organization_id=self.org1_id,
                name="Bob Staff",
                phone="9800000002",
                email="bob@homedesk.com",
                role="organization",
            )
            u3 = User(
                organization_id=self.org2_id,
                name="Charlie Acme",
                phone="9800000003",
                email="charlie@acme.com",
                role="organization",
            )
            db.add_all([u1, u2, u3])
            db.flush()

            org1_users = db.query(User).filter_by(organization_id=self.org1_id).all()
            self.assertGreaterEqual(len(org1_users), 2)
            names_in_org1 = [u.name for u in org1_users]
            self.assertIn("Alice Manager", names_in_org1)
            self.assertIn("Bob Staff", names_in_org1)
            self.assertNotIn("Charlie Acme", names_in_org1)

    def test_02_requirement_belongs_to_category_and_user_not_org(self):
        """Requirements MUST belong to user and category, NOT directly to an organization."""
        mobile = "9123456780"
        with get_db() as db:
            db.query(Requirement).filter_by(mobile=mobile).delete()

        lead_payload = {
            "Name": "Tenant 1 Lead Client",
            "Mobile Number": mobile,
            "City": "Mumbai",
            "Budget (INR)": 16000,
            "Cuisine Type": "North Indian",
            "Meals Per Day": 2,
        }
        saved = save_requirement_to_db("Cook", lead_payload)
        lead_id = saved["Lead ID"]
        self.assertTrue(lead_id.startswith("Cook-"))

        with get_db() as db:
            req = db.query(Requirement).filter_by(lead_id=lead_id).first()
            self.assertIsNotNone(req)
            self.assertEqual(req.category_id, self.org1_cat_id)
            self.assertIsNotNone(req.user_id)
            # Must NOT belong directly to an organization
            self.assertFalse(hasattr(req, "organization_id"))
            self.assertEqual(req.service_type, "Cook")
            self.assertIsNotNone(req.cook_requirement)
            self.assertEqual(req.cook_requirement.cuisine_type, "North Indian")

    def test_03_lead_id_sequence_generation(self):
        """Lead IDs are generated sequentially with standard service prefixes."""
        id1 = generate_next_lead_id("Driver")
        id2 = generate_next_lead_id("Cook")
        self.assertTrue(id1.startswith("Driver-"))
        self.assertTrue(id2.startswith("Cook-"))

    def test_04_cross_category_lookup_by_mobile_prevented(self):
        """Organization A (Cook) can see Cook leads; Organization B (Driver) CANNOT."""
        mobile_a = "9456789012"
        with get_db() as db:
            db.query(Requirement).filter_by(mobile=mobile_a).delete()

        lead_data_a = {
            "Name": "Confidential Cook Lead",
            "Mobile Number": mobile_a,
            "City": "Pune",
            "Budget (INR)": 19000,
            "Cuisine Type": "South Indian",
            "Meals Per Day": 2,
        }
        save_requirement_to_db("Cook", lead_data_a)

        # Query from Org 1 (Cook Organization) finds the lead
        res_org_a = get_latest_lead_by_mobile(mobile_a, organization_id=self.org1_id)
        self.assertIsNotNone(res_org_a)
        self.assertEqual(res_org_a["Name"], "Confidential Cook Lead")

        # Query from Org 2 (Driver Organization) MUST RETURN NONE (Category isolation)
        res_org_b = get_latest_lead_by_mobile(mobile_a, organization_id=self.org2_id)
        self.assertIsNone(
            res_org_b,
            "CRITICAL VIOLATION: Driver organization accessed Cook requirement via mobile lookup!",
        )

    def test_05_cross_category_lookup_by_lead_id_prevented(self):
        """Organization B (Driver) sees Driver leads; Organization A (Cook) cannot."""
        mobile_b = "9567890123"
        with get_db() as db:
            db.query(Requirement).filter_by(mobile=mobile_b).delete()

        lead_data_b = {
            "Name": "Secret Driver Requirement",
            "Mobile Number": mobile_b,
            "City": "Hyderabad",
            "Budget (INR)": 32000,
            "Vehicle Type": "SUV",
            "License Required": "Commercial",
        }
        saved_b = save_requirement_to_db("Driver", lead_data_b)
        lead_id_b = saved_b["Lead ID"]

        # Org 2 (Driver Organization) can view this lead
        found_by_b = get_lead_by_id(lead_id_b, organization_id=self.org2_id)
        self.assertIsNotNone(found_by_b)
        self.assertEqual(found_by_b["Name"], "Secret Driver Requirement")

        # Org 1 (Cook Organization) CANNOT view Driver lead
        found_by_a = get_lead_by_id(lead_id_b, organization_id=self.org1_id)
        self.assertIsNone(
            found_by_a,
            "CRITICAL VIOLATION: Cook organization accessed Driver requirement via Lead ID lookup!",
        )

    def test_06_data_access_layer_enforces_organization_id_validation(self):
        """Data access layer rejects invalid tenant IDs and non-existent organizations."""
        with self.assertRaises(ValueError):
            _resolve_tenant_id(0)

        with self.assertRaises(ValueError):
            _resolve_tenant_id(-5)

        with self.assertRaises(ValueError):
            save_requirement_to_db("Cook", {"Name": "Invalid Org Lead"}, organization_id=999999)

    def test_07_excel_export_strictly_scoped_to_organization_category(self):
        """Organization-scoped Excel export only contains requirements belonging to its category."""
        excel_bytes = get_excel_export_bytes(organization_id=self.org2_id)
        self.assertIsNotNone(excel_bytes)

        wb = load_workbook(io.BytesIO(excel_bytes), data_only=True)
        # Org 2 is Driver: its workbook should only have Driver sheet, no Cook or Security Guard data
        self.assertIn("Driver", wb.sheetnames)
        self.assertNotIn("Cook", wb.sheetnames)
        self.assertNotIn("Security Guard", wb.sheetnames)


if __name__ == "__main__":
    unittest.main(verbosity=2)
